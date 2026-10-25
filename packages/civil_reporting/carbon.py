"""Three export formats from the same immutable carbon project revision."""

from __future__ import annotations

import os
from collections.abc import Mapping
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

import xlsxwriter  # type: ignore[import-untyped]
from civil_api_contracts.carbon import ProjectRevision
from civil_i18n.carbon import carbon_text
from docx import Document
from docx.shared import Pt
from reportlab.lib import colors  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore[import-untyped]
from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
from reportlab.pdfbase.ttfonts import TTFont  # type: ignore[import-untyped]
from reportlab.platypus import (  # type: ignore[import-untyped]
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)


def _fmt(value: float) -> str:
    return f"{value:,.4f}"


def _quantity(value: object) -> str:
    from civil_core.quantity import Quantity

    if not isinstance(value, Quantity):
        raise TypeError("expected Quantity")
    return f"{_fmt(value.input_value)} {value.input_unit}"


def _rows(revision: ProjectRevision, locale: str) -> list[list[str]]:
    def t(key: str) -> str:
        return carbon_text(key, locale)

    result_by_id = {item.row_id: item for item in revision.result.rows}
    rows = [
        [
            t("item"),
            t("category"),
            t("stage"),
            t("quantity"),
            t("factor"),
            t("source"),
            t("evidence"),
            t("emission"),
        ]
    ]
    for item in revision.project.rows:
        factor = item.factor
        emission = result_by_id.get(item.row_id)
        rows.append(
            [
                item.item_name,
                item.category,
                item.stage.value,
                _quantity(item.quantity),
                f"{_quantity(factor.numerator)} / {_quantity(factor.denominator)}"
                if factor else t("pending_factor"),
                f"{factor.source_name} ({factor.source_version}) {factor.source_url or ''}".strip()
                if factor else t("pending_factor"),
                item.evidence_ref or "—",
                f"{_fmt(emission.emission.to('kilogram'))} kgCO2e"
                if emission else t("pending_factor"),
            ]
        )
    return rows


def _metadata(revision: ProjectRevision, locale: str) -> list[tuple[str, str]]:
    def t(key: str) -> str:
        return carbon_text(key, locale)

    lines = [
        (t("project"), revision.project.name),
        (t("asset_type"), revision.project.asset_type.value),
        (t("revision"), str(revision.revision_number)),
        (t("created_at"), revision.created_at.isoformat()),
        (t("hash"), revision.input_snapshot_hash),
        (t("total"), f"{_fmt(revision.result.total_emission.to('kilogram'))} kgCO2e"
         if revision.result.rows else t("pending_factor")),
        (t("reviewer"), t("not_assigned")),
    ]
    if revision.result.pending_row_ids:
        lines.append((t("partial_total"), ", ".join(revision.result.pending_row_ids)))
    if revision.result.tver:
        lines.extend(
            [
                (
                    t("method"),
                    f"{revision.result.tver.methodology_id} "
                    f"v{revision.result.tver.methodology_version}",
                ),
                (t("screening"), revision.result.tver.screening_status),
                (
                    t("estimated_reduction"),
                    f"{_fmt(revision.result.tver.estimated_reduction.to('kilogram'))} kgCO2e",
                ),
                (t("credit_notice"), revision.result.tver.credit_status),
            ]
        )
    return lines


def _tver_details(revision: ProjectRevision, locale: str) -> list[tuple[str, str]]:
    data = revision.project.tver
    if data is None:
        return []
    details = [
        ("BE_y", _quantity(data.baseline_emission)),
        ("PE_y", _quantity(data.project_emission)),
        ("LE_y", _quantity(data.leakage_emission)),
        (carbon_text("component_evidence", locale), data.component_evidence_ref),
        (carbon_text("additionality_evidence", locale), data.additionality_evidence_ref or "—"),
        (carbon_text("credit_rights_evidence", locale), data.credit_rights_evidence_ref or "—"),
        (carbon_text("monitoring_plan", locale), data.monitoring_plan_ref or "—"),
    ]
    if revision.result.tver:
        details.append((carbon_text("findings", locale), ", ".join(revision.result.tver.findings)))
    return details


def _trace_operands(operands: Mapping[str, object]) -> str:
    return "; ".join(f"{key} = {_quantity(value)}" for key, value in operands.items())


def _credit_rows(revision: ProjectRevision, locale: str) -> list[list[str]]:
    names = {row.row_id: row.item_name for row in revision.project.rows}
    rows = [[carbon_text("item", locale), carbon_text("method", locale),
             carbon_text("screening", locale), carbon_text("reference", locale)]]
    for item in revision.result.credit_readiness:
        rows.append([
            names.get(item.row_id, item.row_id),
            f"{item.methodology_id} v{item.methodology_version}" if item.methodology_id else "—",
            (
                f"{carbon_text(item.status, locale)}; "
                f"{carbon_text('credit_' + item.reason_code, locale)}"
            ),
            item.source_url,
        ])
    return rows


def render_carbon_docx(revision: ProjectRevision, locale: str = "th-TH") -> bytes:
    def t(key: str) -> str:
        return carbon_text(key, locale)

    document = Document()
    document.styles["Normal"].font.name = "Tahoma"
    document.styles["Normal"].font.size = Pt(9)
    document.add_heading(t("title"), 0)
    document.add_paragraph(t("draft"))
    for label, value in _metadata(revision, locale):
        document.add_paragraph(f"{label}: {value}")
    for label, value in _tver_details(revision, locale):
        document.add_paragraph(f"{label}: {value}")
    document.add_heading(t("by_stage"), level=1)
    for stage, quantity in revision.result.by_stage.items():
        document.add_paragraph(f"{stage.value}: {_fmt(quantity.to('kilogram'))} kgCO2e")
    document.add_heading(t("activities"), level=1)
    rows = _rows(revision, locale)
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, rows[0], strict=True):
        cell.text = value
    for row in rows[1:]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value
    document.add_heading(t("credit_readiness"), level=1)
    for row in _credit_rows(revision, locale)[1:]:
        document.add_paragraph(" | ".join(row))
    document.add_heading(t("trace"), level=1)
    for item in revision.result.rows:
        trace = item.trace
        document.add_paragraph(
            f"{trace.row_id}: {trace.formula}; {trace.standard_id} "
            f"{trace.edition} §{trace.clause}; {_trace_operands(trace.operands)}; "
            f"{trace.source_url}"
        )
    if revision.result.tver:
        trace = revision.result.tver.trace
        document.add_paragraph(
            f"{trace.formula}; {trace.standard_id} {trace.edition} "
            f"§{trace.clause}; {_trace_operands(trace.operands)}; {trace.source_url}"
        )
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_carbon_xlsx(revision: ProjectRevision, locale: str = "th-TH") -> bytes:
    def t(key: str) -> str:
        return carbon_text(key, locale)

    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    header = workbook.add_format({"bold": True, "bg_color": "#DCEDE8", "font_name": "Tahoma"})
    normal = workbook.add_format({"font_name": "Tahoma"})
    metadata = workbook.add_worksheet("Metadata")
    metadata.set_column(0, 0, 32)
    metadata.set_column(1, 1, 88)
    metadata.write(0, 0, t("title"), header)
    metadata.write(1, 0, t("draft"), normal)
    for index, (label, value) in enumerate(_metadata(revision, locale), start=3):
        metadata.write(index, 0, label, normal)
        metadata.write(index, 1, value, normal)
    if revision.project.tver:
        tver_sheet = workbook.add_worksheet("TVER")
        tver_sheet.set_column(0, 0, 34)
        tver_sheet.set_column(1, 1, 90)
        tver_sheet.write_row(0, 0, [t("method"), revision.project.tver.methodology_id], header)
        for index, (label, value) in enumerate(_tver_details(revision, locale), start=1):
            tver_sheet.write_row(index, 0, [label, value], normal)
    data = workbook.add_worksheet("Activities")
    data.freeze_panes(1, 0)
    data.autofilter(0, 0, len(revision.project.rows), 14)
    data.set_column(0, 1, 24)
    data.set_column(2, 2, 12)
    data.set_column(3, 14, 22)
    headers = [
        t("item"),
        t("category"),
        t("stage"),
        t("quantity"),
        t("quantity_kind"),
        t("unit"),
        t("factor"),
        "kgCO2e",
        t("factor_denominator"),
        t("quantity_kind"),
        t("unit"),
        t("source"),
        t("source_url"),
        t("evidence"),
        t("emission"),
    ]
    data.write_row(0, 0, headers, header)
    by_id = {item.row_id: item for item in revision.result.rows}
    for index, item in enumerate(revision.project.rows, start=1):
        factor = item.factor
        data.write_row(index, 0, [item.item_name, item.category, item.stage.value], normal)
        data.write_number(index, 3, item.quantity.input_value)
        data.write(index, 4, item.quantity.kind)
        data.write(index, 5, item.quantity.input_unit)
        if factor:
            data.write_number(index, 6, factor.numerator.input_value)
            data.write(index, 7, factor.numerator.input_unit)
            data.write_number(index, 8, factor.denominator.input_value)
            data.write(index, 9, factor.denominator.kind)
            data.write(index, 10, factor.denominator.input_unit)
            data.write(index, 11, f"{factor.source_name} ({factor.source_version})")
            data.write(index, 12, factor.source_url or "")
        else:
            data.write(index, 6, t("pending_factor"))
        data.write(index, 13, item.evidence_ref or "")
        if item.row_id in by_id:
            data.write_number(index, 14, by_id[item.row_id].emission.to("kilogram"))
        else:
            data.write(index, 14, t("pending_factor"))
    stages = workbook.add_worksheet("Stages")
    stages.write_row(0, 0, [t("stage"), "kgCO2e"], header)
    for index, (stage, stage_emission) in enumerate(revision.result.by_stage.items(), start=1):
        stages.write(index, 0, stage.value)
        stages.write_number(index, 1, stage_emission.to("kilogram"))
    credit_sheet = workbook.add_worksheet("CreditScreen")
    credit_sheet.set_column(0, 0, 48)
    credit_sheet.set_column(1, 2, 62)
    credit_sheet.set_column(3, 3, 90)
    for index, row in enumerate(_credit_rows(revision, locale)):
        credit_sheet.write_row(index, 0, row, header if index == 0 else normal)
    trace_sheet = workbook.add_worksheet("Trace")
    trace_sheet.set_column(0, 0, 24)
    trace_sheet.set_column(1, 5, 44)
    trace_sheet.write_row(
        0,
        0,
        [
            "Row ID",
            t("formula"),
            t("reference"),
            "Edition",
            t("source_url"),
            "Result kgCO2e",
            t("operands"),
        ],
        header,
    )
    traces = [row.trace for row in revision.result.rows]
    if revision.result.tver:
        traces.append(revision.result.tver.trace)
    for index, step in enumerate(traces, start=1):
        trace_sheet.write_row(
            index,
            0,
            [step.row_id, step.formula, step.standard_id, step.edition, step.source_url],
            normal,
        )
        trace_sheet.write_number(index, 5, step.result.to("kilogram"))
        trace_sheet.write(index, 6, _trace_operands(step.operands))
    workbook.close()
    return output.getvalue()


def _pdf_font(locale: str) -> str:
    if locale == "en-US":
        return "Helvetica"
    configured = os.environ.get("CIVIL_CARBON_PDF_FONT")
    candidates = (
        [Path(configured)]
        if configured
        else [
            Path("C:/Windows/Fonts/tahoma.ttf"),
            Path("/usr/share/fonts/truetype/noto/NotoSansThai-Regular.ttf"),
        ]
    )
    selected = next((path for path in candidates if path.is_file()), None)
    if selected is None:
        raise ValueError("Thai PDF export requires CIVIL_CARBON_PDF_FONT pointing to a Unicode TTF")
    if "CivilCarbonThai" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("CivilCarbonThai", str(selected)))
    return "CivilCarbonThai"


def render_carbon_pdf(revision: ProjectRevision, locale: str = "th-TH") -> bytes:
    def t(key: str) -> str:
        return carbon_text(key, locale)

    font = _pdf_font(locale)
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "CarbonBody", parent=styles["BodyText"], fontName=font, fontSize=9, leading=14
    )
    heading = ParagraphStyle(
        "CarbonHeading", parent=body, fontSize=15, leading=21, textColor=colors.HexColor("#315C55")
    )
    small = ParagraphStyle("CarbonSmall", parent=body, fontSize=7.5, leading=11)
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36
    )
    story: list[object] = [
        Paragraph(escape(t("title")), heading),
        Spacer(1, 8),
        Paragraph(escape(t("draft")), body),
        Spacer(1, 12),
    ]
    for label, value in _metadata(revision, locale):
        story.append(Paragraph(f"<b>{escape(label)}</b>: {escape(value)}", body))
    for label, value in _tver_details(revision, locale):
        story.append(Paragraph(f"<b>{escape(label)}</b>: {escape(value)}", body))
    story.extend([Spacer(1, 12), Paragraph(escape(t("by_stage")), heading)])
    for stage, stage_emission in revision.result.by_stage.items():
        story.append(
            Paragraph(f"{stage.value}: {_fmt(stage_emission.to('kilogram'))} kgCO2e", body)
        )
    story.extend([Spacer(1, 12), Paragraph(escape(t("activities")), heading)])
    for row in _rows(revision, locale)[1:]:
        story.append(Paragraph(" | ".join(escape(cell) for cell in row), small))
        story.append(Spacer(1, 4))
    story.extend([Spacer(1, 12), Paragraph(escape(t("credit_readiness")), heading)])
    for row in _credit_rows(revision, locale)[1:]:
        story.append(Paragraph(" | ".join(escape(cell) for cell in row), small))
        story.append(Spacer(1, 4))
    story.extend([Spacer(1, 12), Paragraph(escape(t("trace")), heading)])
    traces = [row.trace for row in revision.result.rows]
    if revision.result.tver:
        traces.append(revision.result.tver.trace)
    for step in traces:
        story.append(Paragraph(escape(f"{step.row_id}: {step.formula}"), small))
        story.append(Paragraph(escape(_trace_operands(step.operands)), small))
        story.append(
            Paragraph(
                escape(f"{step.standard_id} {step.edition} §{step.clause} — {step.source_url}"),
                small,
            )
        )
    document.build(story)
    return output.getvalue()
